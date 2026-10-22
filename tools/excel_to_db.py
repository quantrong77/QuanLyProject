"""Chuyển file tiến độ Excel (dạng Gantt) sang cơ sở dữ liệu cho web app hiện trường.

Cách dùng:
    python tools/excel_to_db.py "Di Linh Timeline ....xlsx"

Đầu ra:
    data/DiLinh_DB.xlsx      -> upload lên Google Drive, mở bằng Google Sheets làm database
    webapp/demo-data.json    -> dữ liệu cho chế độ DEMO (chạy thử không cần Google)
"""
import json
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
UNIT_RE = re.compile(r"\((SCTNĐ3|NNB|PCC5|PTC3|TOSHIBA)\)", re.I)

# (mã, tên, loại) – danh mục nhà thầu
UNITS = [
    ("SCTNĐ3", "Đội Sửa chữa, thí nghiệm điện 3 (SCTNĐ3) – PTC3", "Thi công"),
    ("NNB", "Đơn vị thi công nhị thứ (NNB)", "Thi công"),
    ("PCC5", "Đơn vị thi công PCCC (PCC5)", "Thi công"),
    ("TOSHIBA", "Toshiba", "Cung cấp thiết bị"),
    ("PTC3", "Công ty Truyền tải điện 3 (PTC3)", "Quản lý dự án"),
    ("PECC4", "Tư vấn giám sát (PECC4)", "Tư vấn giám sát"),
]
UNIT_COLS = ["code", "name", "type", "contact", "phone", "email", "note"]
PROJECT_COLS = ["id", "code", "name", "location", "owner", "start_date", "end_date", "status",
                "units", "description", "created_by", "created_at", "updated_at"]
PROJECT = {
    "id": "P001", "code": "NCS-DILINH", "name": "NCS TBA 500kV Di Linh",
    "location": "TBA 500kV Di Linh, Lâm Đồng", "owner": "Công ty Truyền tải điện 3 (PTC3)",
    "start_date": "2026-10-01", "end_date": "", "status": "Đang thực hiện",
    "units": ",".join(u[0] for u in UNITS),
    "description": "Gồm các hạng mục chính: thay MBA AT1, thay MBA AT2…",
    "created_by": "import", "created_at": "", "updated_at": "",
}
# Hạng mục chính của dự án. File Excel tiến độ hiện có thuộc hạng mục AT1.
PACKAGE_COLS = ["id", "project_id", "code", "name", "start_date", "end_date", "status", "description",
                "order", "created_by", "created_at", "updated_at"]
PACKAGES = [
    {"id": "H001", "project_id": "P001", "code": "AT1", "name": "Thay thế MBA AT1",
     "start_date": "2026-10-01", "end_date": "2026-10-31", "status": "Đang thực hiện",
     "description": "Thay MBA AT1 500/220/35kV 3x150MVA bằng MBA 3x300MVA (Toshiba)", "order": 1, "created_by": "import"},
    {"id": "H002", "project_id": "P001", "code": "AT2", "name": "Thay thế MBA AT2",
     "start_date": "", "end_date": "", "status": "Chuẩn bị",
     "description": "Chưa có tiến độ – nhập đầu việc khi có kế hoạch", "order": 2, "created_by": "import"},
]
IMPORT_PACKAGE = "H001"

TASK_COLS = [
    "id", "project_id", "package_id", "wbs", "section", "group", "subgroup", "name", "phase", "unit",
    "plan_start", "plan_end", "plan_days", "actual_start", "actual_end",
    "progress", "status", "delay_days", "note", "updated_by", "updated_at", "order",
]
UPDATE_COLS = [
    "id", "ts", "report_date", "task_id", "email", "user_name", "unit",
    "progress", "status", "actual_start", "actual_end", "delay_days", "note", "photos",
]
USER_COLS = ["email", "name", "phone", "unit", "role", "status", "created_at", "last_login"]
CONFIG = [
    ("PROJECT_NAME", "NCS TBA 500kV Di Linh"),
    ("ADMIN_EMAILS", "", ),
    ("AUTO_APPROVE", "TRUE"),
    ("ADMIN_UNITS", "PTC3,PECC4"),  # đăng ký với các đơn vị này => quyền Quản trị (chờ duyệt)
    ("GOOGLE_CLIENT_ID", ""),
    ("PHOTO_FOLDER_ID", ""),
]


def iso(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%dT%H:%M") if (v.hour or v.minute) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.isoformat()
    return ""


def days(v):
    if isinstance(v, timedelta):
        return round(v.total_seconds() / 86400, 2)
    if isinstance(v, (int, float)):
        return v
    return ""


def plan_days(start, end, raw):
    """Ưu tiên tính từ ngày bắt đầu/kết thúc (cột thời gian trong Excel bị định dạng giờ)."""
    if isinstance(start, datetime) and isinstance(end, datetime):
        return max(round((end - start).total_seconds() / 86400, 2), 0) or ""
    return days(raw) or ""


# Đơn vị thi công của PTC3: Excel gốc ghi NPTS3, nay do Đội SCTNĐ3 (thuộc PTC3) đảm nhận
RENAME = {"NPTS3": "SCTNĐ3"}


def clean(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip().lstrip("'")
    for old, new in RENAME.items():
        s = s.replace(f"({old})", f"({new})")
    return s


def parse_dm(text, year):
    """'ngày 8/10' -> 2026-10-08"""
    m = re.search(r"(\d{1,2})/(\d{1,2})", text)
    return f"{year}-{int(m.group(2)):02d}-{int(m.group(1)):02d}" if m else ""


def convert(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]

    # Hàng tiêu đề chứa các ngày (cột G trở đi)
    date_cols = {}
    for c in range(7, ws.max_column + 1):
        v = ws.cell(2, c).value
        if isinstance(v, datetime):
            date_cols[c] = v.date()
    year = next(iter(date_cols.values())).year if date_cols else datetime.now().year

    tasks, updates = [], []
    section = group = subgroup = ""
    sec_unit = grp_unit = sub_unit = ""
    sec_idx = -1
    last_name, last_stt = "", ""
    last_num = 0
    wbs_sec = wbs_grp = wbs_sub = ""

    for r in range(4, ws.max_row + 1):
        stt = clean(ws.cell(r, 1).value)
        name = clean(ws.cell(r, 2).value)
        phase_raw = clean(ws.cell(r, 3).value)
        if stt.endswith(".0"):
            stt = stt[:-2]
        if not stt and not name and not phase_raw:
            continue

        # Phân cấp: La Mã tuần tự = phần, chữ hoa = nhóm, chữ thường = nhóm con
        if stt in ROMAN and (sec_idx + 1 < len(ROMAN) and stt == ROMAN[sec_idx + 1]):
            sec_idx += 1
            section, wbs_sec = name, stt
            group = subgroup = wbs_grp = wbs_sub = ""
            last_num = 0
            m = UNIT_RE.search(name)
            sec_unit, grp_unit, sub_unit = (m.group(1).upper() if m else ""), "", ""
            continue
        if re.fullmatch(r"[A-Z]+", stt) and not phase_raw:
            group, wbs_grp = name, stt
            subgroup = wbs_sub = ""
            last_num = 0
            m = UNIT_RE.search(name)
            grp_unit, sub_unit = (m.group(1).upper() if m else ""), ""
            continue
        if re.fullmatch(r"[a-z]", stt):
            subgroup, wbs_sub = name, stt
            last_num = 0
            m = UNIT_RE.search(name)
            sub_unit = m.group(1).upper() if m else ""
            if not ws.cell(r, 5).value:
                continue
            stt = ""
        elif stt.isdigit():
            # Số thứ tự quay lại nhỏ hơn => đã ra khỏi nhóm con
            if subgroup and int(stt) <= last_num:
                subgroup = wbs_sub = sub_unit = ""
            last_num = int(stt)

        if name:
            last_name, last_stt = name, stt
        else:
            name, stt = last_name, last_stt  # hàng pha tiếp theo của cùng công việc

        phase, delay = phase_raw, ""
        m = re.match(r"([^:]*):\s*chậm\s*(\d+)", phase_raw, re.I)
        if m:
            phase, delay = m.group(1).strip(), int(m.group(2))
        if not phase:
            pm = (re.search(r"pha\s+([ABC](?:\s*,\s*[ABC])*)\b", name, re.I)
                  or re.search(r"pha\s+([ABC])\b", subgroup, re.I))
            phase = pm.group(1).upper() if pm else ""

        own = UNIT_RE.search(name)
        unit = (own.group(1).upper() if own else "") or sub_unit or grp_unit or sec_unit
        if not unit and re.search(r"phun sương|PCCC", f"{group} {subgroup} {name}", re.I):
            unit = "PCC5"

        remark = clean(ws.cell(r, 38).value)
        daily = []
        for c, d in date_cols.items():
            v = ws.cell(r, c).value
            # bỏ qua ký hiệu Gantt dạng số (0.5, 1...) – chỉ lấy ghi chú chữ
            if isinstance(v, str) and v.strip() and not re.fullmatch(r"[\d.,\s]+", v.strip()):
                daily.append((d, clean(v)))
        # Ghi chú lặp lại nhiều ngày (VD "Cắt điện") là khung kế hoạch -> đưa vào ghi chú công việc
        repeated = {t for t in {t for _, t in daily} if sum(1 for _, x in daily if x == t) > 1}
        for t in sorted(repeated):
            ds = [d for d, x in daily if x == t]
            remark = (remark + "; " if remark else "") + f"{t}: {ds[0]:%d/%m} → {ds[-1]:%d/%m}"
        daily = [(d, t) for d, t in daily if t not in repeated]

        text = (remark + " " + " ".join(t for _, t in daily)).lower()
        actual_start = actual_end = ""
        progress, status = 0, ""
        if "bắt đầu" in text:
            actual_start = parse_dm(remark.lower().split("bắt đầu")[1], year) if "bắt đầu" in remark.lower() else ""
            if not actual_start:
                actual_start = next((d.isoformat() for d, t in daily if "bắt đầu" in t.lower()), "")
            status = "Đang thi công"
        if "hoàn thành" in text:
            tail = remark.lower().split("hoàn thành")[1] if "hoàn thành" in remark.lower() else ""
            actual_end = parse_dm(tail, year) or next(
                (d.isoformat() for d, t in daily if "hoàn thành" in t.lower()), "")
            progress, status = 100, "Hoàn thành"
        elif status:
            progress = 10

        tid = f"T{len(tasks) + 1:03d}"
        wbs = ".".join(x for x in (wbs_sec, wbs_grp, wbs_sub, stt) if x)
        tasks.append({
            "id": tid, "project_id": PROJECT["id"], "package_id": IMPORT_PACKAGE, "wbs": wbs, "section": f"{wbs_sec}. {section}" if section else "",
            "group": f"{wbs_grp}. {group}" if group else "", "subgroup": subgroup,
            "name": name, "phase": phase, "unit": unit,
            "plan_start": iso(ws.cell(r, 5).value), "plan_end": iso(ws.cell(r, 6).value),
            "plan_days": plan_days(ws.cell(r, 5).value, ws.cell(r, 6).value, ws.cell(r, 4).value),
            "actual_start": actual_start, "actual_end": actual_end,
            "progress": progress, "status": status, "delay_days": delay,
            "note": remark, "updated_by": "import", "updated_at": "", "order": len(tasks) + 1,
        })
        for d, t in daily:
            updates.append({
                "id": f"U{len(updates) + 1:04d}", "ts": f"{d.isoformat()}T17:00",
                "report_date": d.isoformat(), "task_id": tid, "email": "import",
                "user_name": "Nhập từ Excel", "unit": unit, "progress": "", "status": "",
                "actual_start": "", "actual_end": "", "delay_days": "", "note": t, "photos": "",
            })
    return tasks, updates


def write_db(tasks, updates, out):
    wb = openpyxl.Workbook()
    sheets = [
        ("Projects", PROJECT_COLS, [PROJECT]),
        ("Packages", PACKAGE_COLS, PACKAGES),
        ("Tasks", TASK_COLS, tasks),
        ("Updates", UPDATE_COLS, updates),
        ("Users", USER_COLS, []),
        ("Units", UNIT_COLS, units_rows()),
        ("Config", ["key", "value"], [dict(key=k, value=v) for k, v in CONFIG]),
    ]
    wb.remove(wb.active)
    for title, cols, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(cols)
        for row in rows:
            ws.append([row.get(c, "") for c in cols])
        for col in ws.iter_cols(min_row=1, max_row=ws.max_row):
            for cell in col:
                cell.number_format = "@"  # lưu dạng text để Google Sheets không tự đổi ngày
        ws.freeze_panes = "A2"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def units_rows():
    return [dict(code=c, name=n, type=t) for c, n, t in UNITS]


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else next(ROOT.glob("Di Linh Timeline*.xlsx"))
    tasks, updates = convert(src)
    write_db(tasks, updates, ROOT / "data" / "DiLinh_DB.xlsx")
    demo = {"projects": [PROJECT], "packages": PACKAGES, "tasks": tasks, "updates": updates,
            "units": units_rows(),
            "config": {"PROJECT_NAME": CONFIG[0][1], "AUTO_APPROVE": True, "ADMIN_UNITS": ["PTC3", "PECC4"]}}
    (ROOT / "webapp" / "demo-data.json").write_text(json.dumps(demo, ensure_ascii=False, indent=1), encoding="utf-8")
    no_unit = sum(1 for t in tasks if not t["unit"])
    print(f"{len(tasks)} công việc, {len(updates)} ghi chú hiện trường, {no_unit} việc chưa gán đơn vị")


if __name__ == "__main__":
    main()
